"""Crea i file Word dei requisiti (e del business plan) a partire dal Markdown.

Uso:
  python3 tools/requisiti_docx.py REQUISITI.md docs/
      Dal 27/10/2026 (R-511) i requisiti sono DUE Word, generati dallo stesso REQUISITI.md:
        docs/CYF-Gaming-Requisiti-Implementati-e-Testati.docx   ciò che è online e provato
        docs/CYF-Gaming-Requisiti-da-Implementare-o-Bachi.docx  prossimi passi, ciò che manca, scelte e idee
                                                                da confermare, bachi aperti
      Con un'uscita che finisce in .docx (il vecchio uso, un solo Word) i due file vanno nella stessa cartella.
  python3 tools/requisiti_docx.py BUSINESS-PLAN.md uscita.docx    un solo Word, come prima

Come si dividono i requisiti (una riga passa da sola dal secondo al primo Word quando cambia stato):
  - nel primo: stato Fatto, Controllato, Confermata, Superata o Sostituito; le idee "Approvata: ora è R-xxx" se R-xxx
    è fatto; le righe da fare (sezione 9) con ✓ nella colonna Chi; le tabelle della verifica (colonna Esito); le
    tabelle senza stato (per esempio le strade delle città); lo storico fino alla versione online (`online:` in testa
    a REQUISITI.md, la aggiorna chi pubblica); le "Regole di lavoro";
  - nel secondo: tutto il resto (Da fare, In corso, Da confermare, Da approvare...), i "Prossimi passi" in testa e le
    versioni dello storico più nuove di quella online (scritte di giorno, in attesa delle prove della notte).
Una sezione compare in un Word solo se ha almeno una riga per quel Word (o non ha tabelle: allora va nel primo).
Serve pandoc. Le tabelle vengono poi rifinite (bordi, intestazione colorata, colonne).
"""
import os, re, subprocess, sys, tempfile
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NOME_1 = 'CYF-Gaming-Requisiti-Implementati-e-Testati.docx'
NOME_2 = 'CYF-Gaming-Requisiti-da-Implementare-o-Bachi.docx'
TITOLO_1 = 'CYF Gaming Requisiti Implementati e Testati'
TITOLO_2 = 'CYF Gaming Requisiti da Implementare o Bachi'
INTRO_1 = ('**Questo documento contiene solo ciò che è già online e provato**: requisiti Fatto o Controllato, bachi '
           'corretti, scelte confermate, verifiche e storico delle versioni pubblicate. Cosa fare adesso (i prossimi '
           'passi), i requisiti ancora da implementare, le scelte e le idee da confermare e i bachi aperti sono '
           f'nell\'altro documento, «{TITOLO_2}».')
INTRO_2 = ('**Questo documento contiene ciò che resta da fare**: in testa i prossimi passi, poi i bachi aperti, i '
           'requisiti non ancora online (Da fare, In corso), le scelte e le idee da confermare e le versioni scritte '
           'che aspettano le prove della notte. Quando un requisito è online e provato passa da solo nell\'altro '
           f'documento, «{TITOLO_1}».')
FATTO = ('Fatto', 'Controllato', 'Confermata', 'Superata', 'Sostituit')


# ---------------------------------------------------------------- divisione di REQUISITI.md in due Markdown
def celle(riga):
    return [c.strip() for c in re.split(r'(?<!\\)\|', riga.strip())[1:-1]]


def versione(s):
    n = re.findall(r'\d+', s)
    return tuple(int(x) for x in n[:3]) if n else None


def tabelle_e_testo(righe):
    """Divide le righe in pezzi: ('testo', [righe]) oppure ('tabella', intestazione, separatore, [righe])."""
    pezzi, i = [], 0
    while i < len(righe):
        if righe[i].startswith('|') and i + 1 < len(righe) and re.match(r'^\|[\s:|-]+\|\s*$', righe[i + 1]):
            j = i + 2
            while j < len(righe) and righe[j].startswith('|'):
                j += 1
            pezzi.append(('tabella', righe[i], righe[i + 1], righe[i + 2:j]))
            i = j
        else:
            if pezzi and pezzi[-1][0] == 'testo':
                pezzi[-1][1].append(righe[i])
            else:
                pezzi.append(('testo', [righe[i]]))
            i += 1
    return pezzi


def stati_dei_requisiti(righe):
    stati = {}
    for p in tabelle_e_testo(righe):
        if p[0] != 'tabella':
            continue
        h = celle(p[1])
        if 'Stato' not in h:
            continue
        for r in p[3]:
            c = celle(r)
            if len(c) > h.index('Stato'):
                stati[c[0]] = c[h.index('Stato')]
    return stati


def nel_primo(h, c, stati, online):
    """True se la riga va nel Word dei requisiti implementati e testati."""
    if 'Esito' in h:
        return True
    if h and h[0] == 'Versione':
        v, o = versione(c[0]), versione(online or '')
        return not (v and o and v > o)
    if 'Stato' in h and len(c) > h.index('Stato'):
        st = c[h.index('Stato')]
        if st.startswith(FATTO):
            return True
        if st.startswith('Approvat'):
            rif = re.findall(r'\bR-\d+', st) or re.findall(r'\bR-\d+', ' '.join(c))
            return bool(rif) and all(stati.get(x, '').startswith(('Fatto', 'Controllato')) for x in rif)
        return False
    if 'Chi' in h and len(c) > h.index('Chi'):
        return '✓' in c[h.index('Chi')]
    return True


def dividi_blocco(righe, stati, online):
    """Un blocco (una sezione '# ...') → (righe per il primo Word, righe per il secondo); None se non c'è niente."""
    pezzi = tabelle_e_testo(righe)
    ha_tabelle = any(p[0] == 'tabella' for p in pezzi)
    out, conta = ([], []), [0, 0]
    for p in pezzi:
        if p[0] == 'testo':
            out[0].extend(p[1]); out[1].extend(p[1])
            continue
        h = celle(p[1])
        parti = ([], [])
        for r in p[3]:
            parti[0 if nel_primo(h, celle(r), stati, online) else 1].append(r)
        for k in (0, 1):
            if parti[k]:
                out[k].extend([p[1], p[2]] + parti[k]); conta[k] += len(parti[k])
            elif out[k] and out[k][-1].strip():
                out[k].append('')
    if not ha_tabelle:
        return out[0], None
    return (out[0] if conta[0] else None), (out[1] if conta[1] else None)


def dividi(testo):
    righe = testo.split('\n')
    testa = {}
    if righe[0].strip() == '---':
        fine = righe.index('---', 1)
        for r in righe[1:fine]:
            m = re.match(r'^(\w+):\s*"?(.*?)"?\s*$', r)
            if m:
                testa[m.group(1)] = m.group(2)
        righe = righe[fine + 1:]
    online = testa.get('online', '')
    stati = stati_dei_requisiti(righe)
    # blocchi di primo livello
    blocchi, cur = [], None
    for r in righe:
        if r.startswith('# '):
            cur = [r]; blocchi.append(cur)
        elif cur is not None:
            cur.append(r)
    doc1, doc2, regole, leggere = [], [], [], []
    for b in blocchi:
        titolo = b[0][2:].strip()
        if titolo.startswith('Prossimi passi'):
            k = next((i for i, r in enumerate(b) if r.startswith('## Regole di lavoro')), len(b))
            doc2.extend(b[:k])
            if k < len(b):
                regole = ['# ' + b[k][3:]] + b[k + 1:]
            continue
        if titolo.startswith('Come leggere'):
            leggere = b
            continue
        p1, p2 = dividi_blocco(b, stati, online)
        if titolo.startswith('Storico delle versioni') and p2:
            p2 = ['# Versioni scritte, in attesa delle prove della notte', '',
                  f'Più nuove della versione online ({online}): sono sui rami `giorno/…`; il turno di notte le prova '
                  f'e le pubblica, poi passano nell\'altro documento.', ''] + p2[1:]
        if p1:
            if titolo.endswith('(da confermare)'):
                p1 = [p1[0].replace(' (da confermare)', '')] + p1[1:]
            if titolo == '9. Da fare':
                p1 = ['# 9. Cose da fare già fatte'] + p1[1:]
            doc1.extend(p1)
        if p2:
            doc2.extend(p2)
    sub = testa.get('subtitle', '')
    data = re.search(r'aggiornato al [^·]+', sub)
    data = data.group(0).strip() if data else ''

    def documento(titolo, sottotitolo, intro, corpo, prima):
        come = []
        if leggere:
            come = [leggere[0], '', intro, ''] + leggere[1:]
        testa_md = ['---', f'title: "{titolo}"', f'subtitle: "{sottotitolo}"', '---', '']
        return '\n'.join(testa_md + (prima + come + corpo if prima is not None else corpo + come)) + '\n'

    sotto1 = ' · '.join(x for x in (f'Versione online {online}' if online else '', data,
                                    'solo ciò che è online e provato') if x)
    md1 = documento(TITOLO_1, sotto1, INTRO_1, regole + doc1, [])
    # nel secondo Word i prossimi passi restano in testa, poi "come leggere", poi il resto
    k = next((i for i, r in enumerate(doc2[1:], 1) if r.startswith('# ')), len(doc2))
    md2 = documento(TITOLO_2, sub, INTRO_2, doc2[k:], doc2[:k])
    return md1, md2


# ---------------------------------------------------------------- impaginazione del Word (come prima)
STATE = {'Controllato': '9ED9B3', 'Confermata 25/10': '9ED9B3', 'Superata 25/10': 'ECECEC', 'Sostituito': 'ECECEC', 'Fatto': 'D7F0DF', 'In corso': 'FFF1C7', 'Da fare': 'FDE2D6', 'Da confermare': 'E6EDF5', '✓': 'D7F0DF', '✓ dopo correzione': 'FFF1C7', 'Sì': 'D7F0DF', 'Da approvare': 'FFF1C7'}


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'B8C4D2')
        b.append(e)
    tblPr.append(b)


def word(src_md, out, BP):
    subprocess.run(['pandoc', src_md, '-f', 'markdown', '-t', 'docx', '-o', out], check=True)
    doc = Document(out)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin = s.bottom_margin = Cm(1.8)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    for st in doc.styles:
        if st.name in ('Title', 'Subtitle') or st.name.lower().startswith('heading'):
            st.font.name = 'Calibri'
            st.font.color.rgb = RGBColor(0x13, 0x23, 0x3A)
    names = {st.name for st in doc.styles}
    if 'Title' in names: doc.styles['Title'].font.size = Pt(22)
    for st in doc.styles:
        if st.name.lower().startswith('heading 1'):
            st.font.size = Pt(15)
    WIDTHS = {5: [1.9, 4.3, 3.2, 3.2, 4.8], 4: [2.3, 3.4, 7.2, 4.5], 3: [4.3, 5.4, 7.7], 2: [4.9, 12.5]} if BP else {5: [1.6, 7.4, 2.0, 2.0, 4.4], 4: [1.4, 10.6, 2.2, 3.2], 3: [1.6, 12.0, 3.8], 6: [0.9, 5.6, 4.4, 2.2, 2.0, 2.3]}
    for t in doc.tables:
        borders(t)
        n = len(t.columns)
        w = WIDTHS.get(n)
        if w and n == 3 and t.rows and t.rows[0].cells[0].text.strip() == 'Versione':
            w = [2.2, 2.4, 12.8]   # lo storico ha la descrizione nell'ultima colonna
        if w and n == 5 and t.rows and t.rows[0].cells[0].text.strip() == 'Città':
            w = [1.7, 4.2, 4.3, 4.3, 3.9]   # le strade delle città
        if w and n == 5 and len(t.rows) > 1 and t.rows[1].cells[0].text.strip().startswith('A-'):
            w = [1.4, 4.8, 2.0, 1.8, 7.4]   # struttura del software: spiegazioni lunghe nell'ultima colonna
        if w and n == 5 and len(t.rows) > 1 and t.rows[1].cells[0].text.strip().startswith('B-'):
            w = [1.4, 6.0, 2.0, 1.8, 6.2]   # bachi: causa e correzione nell'ultima colonna
        if w and n == 4 and t.rows and t.rows[0].cells[0].text.strip() == 'N.':
            w = [0.9, 11.6, 2.2, 2.7]   # prossimi passi di Luigi: numero, cosa, tempo, ID
        if w and n == 3 and t.rows and t.rows[0].cells[1].text.strip() == 'Esito':
            w = [3.4, 2.6, 11.4]   # la verifica: ID, esito, come è stato verificato
        if w:
            # pandoc crea colonne uguali: si fissano larghezze e griglia della tabella
            t.autofit = False
            tblPr = t._tbl.tblPr
            for old in tblPr.findall(qn('w:tblW')) + tblPr.findall(qn('w:tblLayout')):
                tblPr.remove(old)
            tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(w) * 567))); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
            lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
            for gc, cw in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), w):
                gc.set(qn('w:w'), str(int(cw * 567)))
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                if w:
                    cell.width = Cm(w[ci])
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        if ri == 0:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                if ri == 0:
                    shade(cell, '13233A')
                else:
                    txt = cell.text.strip()
                    if txt in STATE:
                        shade(cell, STATE[txt])
    doc.save(out)
    print('ok', out)


def main():
    src, out = sys.argv[1], sys.argv[2]
    BP = os.path.basename(src).lower().startswith('business')
    if BP:
        word(src, out, True)
        return
    cartella = os.path.dirname(out) if out.endswith('.docx') else out
    os.makedirs(cartella or '.', exist_ok=True)
    md1, md2 = dividi(open(src, encoding='utf-8').read())
    with tempfile.TemporaryDirectory() as tmp:
        for md, nome in ((md1, NOME_1), (md2, NOME_2)):
            p = os.path.join(tmp, nome.replace('.docx', '.md'))
            open(p, 'w', encoding='utf-8').write(md)
            word(p, os.path.join(cartella, nome), False)


if __name__ == '__main__':
    main()
